#!/usr/bin/env python3
"""Independent integrity checks for the generated Rwanda landslide v2 tables."""
from pathlib import Path
import json, re, sys, xml.etree.ElementTree as ET
import pandas as pd
import geopandas as gpd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
errors=[]
def need(ok,msg):
    if not ok: errors.append(msg)

def read(name):
    p=ROOT/name
    need(p.exists(),f"missing {name}")
    return pd.read_csv(p,keep_default_na=False) if p.exists() else pd.DataFrame()

E=read('events.csv');S=read('sites.csv');SRC=read('sources.csv');V=read('site_validation.csv');X=read('source_event_crosswalk.csv')
allowed_types={'shallow landslide','debris flow','mudslide','rockfall','rotational/translational slide','landslide cluster','slope failure (mining/road/construction)','mixed landslide-flood','unknown'}
allowed_triggers={'rainfall','prolonged rainfall','earthquake','mining/quarrying','construction/road cut','flood-related','unknown'}
need(len(E)>0 and E.event_id.is_unique,'event_id must be unique/nonempty')
need(len(S)>0 and S.site_id.is_unique,'site_id must be unique/nonempty')
need(set(S.event_id)<=set(E.event_id),'site has missing event foreign key')
need(set(E.landslide_type)<=allowed_types,'invalid event landslide_type')
need(set(E.primary_trigger)<=allowed_triggers,'invalid event primary_trigger')
need(set(E.event_confidence)<={'High','Medium','Low'},'invalid event confidence')
need(set(E.date_precision)<={'day','month','year','range'},'invalid date precision')
need(set(S.geometry_type)<={'point','polygon','centroid_of_polygon','aoi_center'},'invalid geometry_type')
need(set(S.location_quality)<={'exact','approximate','coarse','aoi_center'},'invalid location_quality')
need(set(S.coordinate_method)<={'published','inventory','UNOSAT','geocoded','derived'},'invalid coordinate_method')
for c in ['date_start']:
    parsed=pd.to_datetime(E[c],errors='coerce',format='%Y-%m-%d')
    need(parsed.notna().all(),f'invalid {c}')
end=E.date_end[E.date_end!=''];need(pd.to_datetime(end,errors='coerce',format='%Y-%m-%d').notna().all(),'invalid date_end')
for c in ['n_districts','fatalities_min','fatalities_max','injuries_min','injuries_max','people_affected','people_displaced','houses_destroyed','houses_damaged']:
    if c in E:
        bad=E[c].astype(str).ne('') & ~E[c].astype(str).str.fullmatch(r'\d+')
        need(not bad.any(),f'non-numeric event field {c}')
need(not (E.fatalities_min!='').any() or (E[E.fatalities_min!=''].fatalities_min.astype(int)<=E[E.fatalities_min!=''].fatalities_max.astype(int)).all(),'fatality min exceeds max')
source_ids=set(SRC.source_id);need(SRC.source_id.is_unique,'source_id must be unique')
need(SRC.url.astype(str).str.startswith(('http://','https://')).all(),'source URL missing/invalid')
for _,e in E.iterrows():
    need(bool(e.source_ids),f"event has no sources: {e.event_id}")
    need(set(filter(None,e.source_ids.split('|'))) <= source_ids,f"event source FK missing: {e.event_id}")
for _,s in S.iterrows():
    need(bool(s.source_id) and s.source_id in source_ids,f"site source FK missing: {s.site_id}")
    haslat=s.latitude!=''; haslon=s.longitude!='';need(haslat==haslon,f"partial coordinate pair: {s.site_id}")
    if haslat:
        lat=float(s.latitude);lon=float(s.longitude)
        need(-2.85<=lat<=-1.04 and 28.85<=lon<=30.91,f"coordinate outside Rwanda envelope: {s.site_id}")
need(len(V)==len(S) and set(V.site_id)==set(S.site_id),'site validation coverage must match sites')
need(set(V.district_check)<={'pass','fail','not_tested'},'invalid district QA status')
need(set(V.sector_check)<={'pass','fail','not_tested'},'invalid sector QA status')
need(not ((V.district_check=='fail') & V.mismatch_explanation.eq('')).any(),'unexplained district failure')
need(not ((V.sector_check=='fail') & V.mismatch_explanation.eq('')).any(),'unexplained sector failure')
need(len(X)>0 and X.loc[X.event_id!='','event_id'].isin(E.event_id).all(),'source crosswalk has orphan event ID')
# Spatial and office formats.
for name,layer,n in [('sites.geojson',None,len(S)),('rwanda_landslide_v2.gpkg','sites',len(S)),('rwanda_landslide_v2.gpkg','events',len(E))]:
    try:
        d=gpd.read_file(ROOT/name,layer=layer) if layer else gpd.read_file(ROOT/name)
        need(len(d)==n,f'{name}/{layer} row count mismatch')
        need(str(d.crs).endswith('4326'),f'{name}/{layer} must use EPSG:4326')
    except Exception as ex:errors.append(f'cannot read {name}/{layer}: {ex}')
try:
    with (ROOT/'events.geojson').open(encoding='utf-8') as f:g=json.load(f)
    need(len(g.get('features',[]))==len(E),'events GeoJSON feature count mismatch')
except Exception as ex:errors.append(f'cannot read events.geojson: {ex}')
try:
    ET.parse(ROOT/'sites.kml')
except Exception as ex:errors.append(f'cannot parse sites.kml: {ex}')
try:
    w=load_workbook(ROOT/'rwanda_landslide_v2.xlsx',read_only=True,data_only=True)
    need({'Events','Sites','Sources','README'}<=set(w.sheetnames),'workbook is missing a required core/README sheet')
    need(w['Events'].max_row==len(E)+1,'workbook Events row count mismatch')
    need(w['Sites'].max_row==len(S)+1,'workbook Sites row count mismatch')
except Exception as ex:errors.append(f'cannot read workbook: {ex}')
if errors:
    print('VALIDATION FAILED')
    for e in errors:print('- '+e)
    sys.exit(1)
print(json.dumps({'errors':0,'events':len(E),'sites':len(S),'coordinate_pairs':int(S.latitude.ne('').sum()),'sources':len(SRC),'district_failures':int((V.district_check=='fail').sum()),'sector_failures':int((V.sector_check=='fail').sum()),'crosswalk_rows':len(X)},indent=2))
